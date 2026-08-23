"""文档解析器（第三十二节）。

支持 PDF（PyMuPDF/fitz）、DOCX（python-docx）、TXT/MD。
解析 → 表格结构化 → 清洗 → 返回纯文本。不连真实设备/不调用外部服务。
未安装解析依赖时优雅降级并明确报错。

表格策略：PDF 用 find_tables 识别并以 Markdown 表格文本输出，
同时从页面文本中剔除表格区域避免重复；DOCX 按文档元素顺序
交替输出段落与表格（paragraphs 不含表格，单独遍历会丢表）。
页眉页脚：跨页频次启发，高频短行视为页眉/页脚剔除。
"""

from __future__ import annotations

import re
import zipfile
from collections import Counter
from dataclasses import dataclass, field
from html.parser import HTMLParser
from pathlib import Path

from app.rag.cleaner import clean_markdown, clean_text

_TABLE_MARKER = "【表格】"
# 页眉/页脚判定：短行且出现在超过该比例页面中
_REPEATED_LINE_RATIO = 0.4
_REPEATED_LINE_MAX_LEN = 40
_HTML_TABLE = re.compile(r"<table\b[^>]*>.*?</table\s*>", re.IGNORECASE | re.DOTALL)
_FENCE_LINE = re.compile(r"^\s*(```|~~~)")
_MAX_HTML_TABLE_SPAN = 64


@dataclass
class ParsedDocument:
    """解析结果。"""

    text: str
    source_path: str
    source_type: str  # pdf/docx/txt/md
    page_count: int = 0
    error: str = ""
    pages: list[str] = field(default_factory=list)


@dataclass(slots=True)
class _HtmlTableCell:
    text: str
    rowspan: int = 1
    colspan: int = 1


class _HtmlTableParser(HTMLParser):
    """只提取表格的文本和跨行/跨列信息，不保留任意原始 HTML。"""

    def __init__(self) -> None:
        super().__init__(convert_charrefs=True)
        self.rows: list[list[_HtmlTableCell]] = []
        self._row: list[_HtmlTableCell] | None = None
        self._cell_parts: list[str] | None = None
        self._cell_rowspan = 1
        self._cell_colspan = 1
        self._ignored_depth = 0

    @staticmethod
    def _span(attrs: dict[str, str | None], name: str) -> int:
        try:
            value = int(attrs.get(name) or "1")
        except ValueError:
            return 1
        return min(max(value, 1), _MAX_HTML_TABLE_SPAN)

    def handle_starttag(self, tag: str, attrs: list[tuple[str, str | None]]) -> None:
        lower = tag.lower()
        if lower in {"script", "style"}:
            self._ignored_depth += 1
            return
        if self._ignored_depth:
            return
        if lower == "tr":
            self._row = []
        elif lower in {"td", "th"} and self._row is not None:
            attr_map = {key.lower(): value for key, value in attrs}
            self._cell_parts = []
            self._cell_rowspan = self._span(attr_map, "rowspan")
            self._cell_colspan = self._span(attr_map, "colspan")
        elif lower == "br" and self._cell_parts is not None:
            self._cell_parts.append(" ")

    def handle_endtag(self, tag: str) -> None:
        lower = tag.lower()
        if lower in {"script", "style"} and self._ignored_depth:
            self._ignored_depth -= 1
            return
        if self._ignored_depth:
            return
        if lower in {"td", "th"} and self._cell_parts is not None and self._row is not None:
            text = " ".join("".join(self._cell_parts).split())
            self._row.append(_HtmlTableCell(text, self._cell_rowspan, self._cell_colspan))
            self._cell_parts = None
        elif lower == "tr" and self._row is not None:
            if self._row:
                self.rows.append(self._row)
            self._row = None

    def handle_data(self, data: str) -> None:
        if self._cell_parts is not None and not self._ignored_depth:
            self._cell_parts.append(data)


def parse_document(
    path: str | Path,
    *,
    max_pages: int | None = None,
    max_docx_entries: int | None = None,
    max_docx_uncompressed_bytes: int | None = None,
    max_docx_compression_ratio: int | None = None,
) -> ParsedDocument:
    """按扩展名分发解析；PDF 可在读取前限制页数以控制内存峰值。"""
    p = Path(path)
    if not p.exists():
        return ParsedDocument(text="", source_path=str(p), source_type="", error="文件不存在")
    ext = p.suffix.lower().lstrip(".")
    if ext == "pdf":
        return _parse_pdf(p, max_pages=max_pages)
    if ext == "docx":
        return _parse_docx(
            p,
            max_entries=max_docx_entries,
            max_uncompressed_bytes=max_docx_uncompressed_bytes,
            max_compression_ratio=max_docx_compression_ratio,
        )
    if ext in ("txt", "md", "markdown"):
        return _parse_text(p)
    return ParsedDocument(text="", source_path=str(p), source_type=ext, error=f"不支持的格式: {ext}")


def _render_table(rows: list[list]) -> str:
    """将表格行渲染为 Markdown 表格文本（保留结构，供分块与检索）。"""
    clean_rows: list[list[str]] = []
    for row in rows:
        cells = [("" if cell is None else str(cell)).replace("\n", " ").strip() for cell in row]
        if any(cells):
            clean_rows.append(cells)
    if not clean_rows:
        return ""
    width = max(len(r) for r in clean_rows)
    lines = [_TABLE_MARKER]
    for index, row in enumerate(clean_rows):
        padded = row + [""] * (width - len(row))
        lines.append("| " + " | ".join(padded) + " |")
        if index == 0:
            lines.append("|" + " --- |" * width)
    return "\n".join(lines)


def _expand_html_table_rows(rows: list[list[_HtmlTableCell]]) -> list[list[str]]:
    """将 rowspan/colspan 展开为规则矩形，GFM 不支持合并单元格。"""

    active_spans: dict[int, int] = {}
    expanded: list[list[str]] = []
    width = 0
    for source_row in rows:
        cells: dict[int, str] = {}
        column = 0

        def consume_active_spans(current_cells: dict[int, str]) -> None:
            nonlocal column
            while column in active_spans:
                current_cells[column] = ""
                active_spans[column] -= 1
                if active_spans[column] <= 0:
                    del active_spans[column]
                column += 1

        for cell in source_row:
            consume_active_spans(cells)
            for offset in range(cell.colspan):
                target = column + offset
                cells[target] = cell.text if offset == 0 else ""
                if cell.rowspan > 1:
                    active_spans[target] = max(active_spans.get(target, 0), cell.rowspan - 1)
            column += cell.colspan
        while active_spans:
            next_column = min(active_spans)
            if next_column < column:
                break
            column = next_column
            consume_active_spans(cells)
        row_width = max(cells, default=-1) + 1
        width = max(width, row_width)
        expanded.append([cells.get(index, "") for index in range(row_width)])
    return [row + [""] * (width - len(row)) for row in expanded]


def _html_table_to_markdown(table_html: str) -> str:
    parser = _HtmlTableParser()
    try:
        parser.feed(table_html)
        parser.close()
    except Exception:  # noqa: BLE001 - 解析异常时保留原文，避免丢失资料
        return table_html
    rows = _expand_html_table_rows(parser.rows)
    if not rows or not any(any(cell for cell in row) for row in rows):
        return table_html
    return _render_table(rows)


def normalize_html_tables(text: str) -> str:
    """把 MinerU 等工具输出的 HTML 表格转为安全、可分块的 GFM 表格。

    围栏代码中的示例 HTML 不转换，避免改写用户的代码片段。
    """

    if "<table" not in text.lower():
        return text
    parts: list[str] = []
    buffer: list[str] = []
    in_fence = False

    def flush_buffer() -> None:
        if buffer:
            parts.append(_HTML_TABLE.sub(lambda match: _html_table_to_markdown(match.group(0)), "".join(buffer)))
            buffer.clear()

    for line in text.splitlines(keepends=True):
        if _FENCE_LINE.match(line):
            flush_buffer()
            parts.append(line)
            in_fence = not in_fence
        elif in_fence:
            parts.append(line)
        else:
            buffer.append(line)
    flush_buffer()
    return "".join(parts)


def _rect_overlap_ratio(block: tuple, rect) -> float:
    """文本块与表格区域的重叠面积占块面积比例。"""
    x0, y0, x1, y1 = block[0], block[1], block[2], block[3]
    ix0, iy0 = max(x0, rect.x0), max(y0, rect.y0)
    ix1, iy1 = min(x1, rect.x1), min(y1, rect.y1)
    inter = max(0.0, ix1 - ix0) * max(0.0, iy1 - iy0)
    area = max(1e-6, (x1 - x0) * (y1 - y0))
    return inter / area


def _remove_repeated_lines(pages: list[str]) -> list[str]:
    """跨页频次启发式去除页眉/页脚：同一短行出现在过多页面。"""
    if len(pages) < 3:
        return pages
    counter: Counter[str] = Counter()
    for page in pages:
        seen_in_page: set[str] = set()
        for line in page.splitlines():
            stripped = line.strip()
            if 0 < len(stripped) <= _REPEATED_LINE_MAX_LEN and stripped not in seen_in_page:
                seen_in_page.add(stripped)
                counter[stripped] += 1
    threshold = len(pages) * _REPEATED_LINE_RATIO
    noise = {line for line, count in counter.items() if count >= threshold}
    if not noise:
        return pages
    result: list[str] = []
    for page in pages:
        kept = [line for line in page.splitlines() if line.strip() not in noise]
        result.append("\n".join(kept))
    return result


def _parse_pdf(p: Path, *, max_pages: int | None = None) -> ParsedDocument:
    doc = None
    try:
        import fitz  # PyMuPDF
    except Exception as exc:  # noqa: BLE001
        return ParsedDocument(text="", source_path=str(p), source_type="pdf", error=f"PyMuPDF 不可用: {exc}")
    try:
        doc = fitz.open(str(p))
        page_count = len(doc)
        if max_pages is not None and page_count > max_pages:
            return ParsedDocument(
                text="",
                source_path=str(p),
                source_type="pdf",
                page_count=page_count,
                error=f"文档页数超限（{page_count} 页），最大 {max_pages} 页",
            )
        parts: list[str] = []
        for page in doc:
            # 1. 表格识别：结构化输出并从正文中剔除表格区域，避免重复
            table_texts: list[str] = []
            table_rects: list = []
            try:
                for table in page.find_tables().tables:
                    rendered = _render_table(table.extract())
                    if rendered:
                        table_texts.append(rendered)
                        table_rects.append(fitz.Rect(table.bbox))
            except Exception:  # noqa: BLE001 - find_tables 不可用时降级纯文本
                pass
            text_parts: list[str] = []
            try:
                blocks = page.get_text("blocks")
                for block in blocks:
                    if len(block) < 7 or block[6] != 0:  # 仅文本块
                        continue
                    if any(_rect_overlap_ratio(block, rect) > 0.5 for rect in table_rects):
                        continue
                    text_parts.append(str(block[4]))
            except Exception:  # noqa: BLE001
                text_parts.append(page.get_text("text"))
            page_text = "\n".join(text_parts)
            if table_texts:
                page_text = page_text + "\n" + "\n".join(table_texts)
            parts.append(clean_text(page_text))
        parts = _remove_repeated_lines(parts)
        text = clean_text("\n\n".join(parts))
        return ParsedDocument(
            text=text,
            source_path=str(p),
            source_type="pdf",
            page_count=page_count,
            pages=parts,
        )
    except Exception as exc:  # noqa: BLE001
        return ParsedDocument(text="", source_path=str(p), source_type="pdf", error=str(exc))
    finally:
        if doc is not None:
            doc.close()


def _validate_docx_archive(
    p: Path,
    *,
    max_entries: int | None,
    max_uncompressed_bytes: int | None,
    max_compression_ratio: int | None,
) -> str:
    """在解压前拒绝异常 DOCX，避免小 ZIP 在低内存容器中膨胀。"""

    try:
        with zipfile.ZipFile(p) as archive:
            entries = [entry for entry in archive.infolist() if not entry.is_dir()]
    except (OSError, zipfile.BadZipFile) as exc:
        return f"DOCX 压缩包无效: {exc}"
    if max_entries is not None and len(entries) > max_entries:
        return f"DOCX 文件条目超限（{len(entries)} 个），最大 {max_entries} 个"
    uncompressed_bytes = sum(entry.file_size for entry in entries)
    if max_uncompressed_bytes is not None and uncompressed_bytes > max_uncompressed_bytes:
        return (
            f"DOCX 解压后体积超限（{uncompressed_bytes} 字节），"
            f"最大 {max_uncompressed_bytes} 字节"
        )
    if max_compression_ratio is not None:
        for entry in entries:
            if entry.file_size and entry.compress_size == 0:
                return f"DOCX 压缩比异常: {entry.filename}"
            if entry.compress_size and entry.file_size / entry.compress_size > max_compression_ratio:
                return f"DOCX 压缩比超限: {entry.filename}"
    return ""


def _parse_docx(
    p: Path,
    *,
    max_entries: int | None = None,
    max_uncompressed_bytes: int | None = None,
    max_compression_ratio: int | None = None,
) -> ParsedDocument:
    archive_error = _validate_docx_archive(
        p,
        max_entries=max_entries,
        max_uncompressed_bytes=max_uncompressed_bytes,
        max_compression_ratio=max_compression_ratio,
    )
    if archive_error:
        return ParsedDocument(text="", source_path=str(p), source_type="docx", error=archive_error)
    try:
        import docx  # python-docx
        from docx.oxml.ns import qn
        from docx.table import Table as DocxTable
        from docx.text.paragraph import Paragraph
    except Exception as exc:  # noqa: BLE001
        return ParsedDocument(text="", source_path=str(p), source_type="docx", error=f"python-docx 不可用: {exc}")
    try:
        d = docx.Document(str(p))
        parts: list[str] = []

        def append_block(value: str) -> None:
            if value:
                parts.extend((value, ""))

        def heading_prefix(style_name: str) -> str | None:
            lowered = style_name.lower()
            if not (lowered.startswith("heading") or "标题" in style_name):
                return None
            match = re.search(r"(\d+)", style_name)
            level = min(max(int(match.group(1)) if match else 1, 1), 6)
            return "#" * level

        def list_prefix(para: Paragraph, style_name: str) -> str:
            ppr = para._p.pPr
            num_pr = getattr(ppr, "numPr", None) if ppr is not None else None
            if num_pr is None:
                lowered = style_name.lower()
                if "list bullet" in lowered or "项目符号" in style_name:
                    return "- "
                if "list number" in lowered or "编号" in style_name:
                    return "1. "
                return ""
            level_node = getattr(num_pr, "ilvl", None)
            level = int(level_node.val) if level_node is not None and level_node.val else 0
            return f"{'  ' * level}- "

        # 按文档体元素顺序交替处理段落与表格，保留原始位置关系
        for child in d.element.body.iterchildren():
            if child.tag == qn("w:p"):
                para = Paragraph(child, d)
                value = para.text.strip()
                if not value:
                    continue
                style_name = getattr(para.style, "name", "") or ""
                prefix = heading_prefix(style_name)
                append_block(f"{prefix} {value}" if prefix else f"{list_prefix(para, style_name)}{value}")
            elif child.tag == qn("w:tbl"):
                table = DocxTable(child, d)
                rows = [[cell.text.strip() for cell in row.cells] for row in table.rows]
                rendered = _render_table(rows)
                if rendered:
                    append_block(rendered)
        text = clean_markdown("\n".join(parts))
        return ParsedDocument(
            text=text,
            source_path=str(p),
            source_type="docx",
            page_count=1,
            pages=[text],
        )
    except Exception as exc:  # noqa: BLE001
        return ParsedDocument(text="", source_path=str(p), source_type="docx", error=str(exc))


def _parse_text(p: Path) -> ParsedDocument:
    try:
        raw = p.read_text(encoding="utf-8", errors="replace")
    except Exception as exc:  # noqa: BLE001
        return ParsedDocument(text="", source_path=str(p), source_type="txt", error=str(exc))
    ext = p.suffix.lower().lstrip(".")
    if ext in ("md", "markdown"):
        text = clean_markdown(normalize_html_tables(raw))
    else:
        text = clean_text(raw)
    return ParsedDocument(
        text=text,
        source_path=str(p),
        source_type=ext,
        page_count=1,
        pages=[text],
    )


__all__ = ["ParsedDocument", "normalize_html_tables", "parse_document"]
