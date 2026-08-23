"""文档解析的格式保真与资源边界测试。"""

import zipfile
from pathlib import Path

from app.rag.parser import normalize_html_tables, parse_document
from app.rag.structured_chunker import split_structured_text


def test_docx_parser_emits_headings_lists_and_table_markdown(tmp_path: Path):
    import docx

    document = docx.Document()
    document.add_heading("阀门巡检", level=1)
    document.add_paragraph("确认阀门编号", style="List Number")
    document.add_paragraph("检查密封填料", style="List Bullet")
    table = document.add_table(rows=2, cols=2)
    table.cell(0, 0).text = "项目"
    table.cell(0, 1).text = "要求"
    table.cell(1, 0).text = "球阀"
    table.cell(1, 1).text = "无渗漏"
    path = tmp_path / "valve.docx"
    document.save(path)

    parsed = parse_document(path)

    assert not parsed.error
    assert "# 阀门巡检" in parsed.text
    assert "1. 确认阀门编号" in parsed.text
    assert "- 检查密封填料" in parsed.text
    assert "| 球阀 | 无渗漏 |" in parsed.text


def test_pdf_page_limit_is_checked_before_text_extraction(tmp_path: Path):
    import fitz

    path = tmp_path / "two-pages.pdf"
    document = fitz.open()
    document.new_page().insert_text((72, 72), "第一页")
    document.new_page().insert_text((72, 72), "第二页")
    document.save(path)
    document.close()

    parsed = parse_document(path, max_pages=1)

    assert parsed.error == "文档页数超限（2 页），最大 1 页"
    assert parsed.page_count == 2
    assert not parsed.text


def test_docx_archive_is_rejected_before_unzip_when_it_exceeds_limit(tmp_path: Path):
    path = tmp_path / "oversized.docx"
    with zipfile.ZipFile(path, "w", compression=zipfile.ZIP_DEFLATED) as archive:
        archive.writestr("word/document.xml", "x" * 1_024)

    parsed = parse_document(path, max_docx_uncompressed_bytes=100)

    assert "DOCX 解压后体积超限" in parsed.error


def test_mineru_html_table_is_normalized_before_structured_chunking(tmp_path: Path):
    path = tmp_path / "mineru.md"
    path.write_text(
        """# 燃气输配设施

<table><tr><td>职业功能</td><td>工作内容</td></tr><tr><td rowspan="2">燃气输配</td><td>压力测试</td></tr><tr><td>置换通气<script>alert(1)</script></td></tr></table>
""",
        encoding="utf-8",
    )

    parsed = parse_document(path)
    chunks = split_structured_text(parsed.text)

    assert "<table" not in parsed.text
    assert "alert(1)" not in parsed.text
    assert "| 职业功能 | 工作内容 |" in parsed.text
    assert "| 燃气输配 | 压力测试 |" in parsed.text
    assert "|  | 置换通气 |" in parsed.text
    assert len(chunks) == 1
    assert chunks[0].chunk_type == "table"


def test_html_table_inside_code_fence_is_not_normalized():
    text = "```html\n<table><tr><td>示例</td></tr></table>\n```"

    assert normalize_html_tables(text) == text
